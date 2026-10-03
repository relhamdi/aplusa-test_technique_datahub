import codecs
import csv
from collections.abc import Iterator
from dataclasses import dataclass
from typing import BinaryIO

import pandas as pd
from openpyxl import load_workbook

SAMPLE_BYTES = 64 * 1024
CANDIDATE_DELIMITERS = ",;\t|"


class FileReadError(Exception):
    """Raised for any unreadable or unsupported upload (mapped to HTTP 400)."""


@dataclass(frozen=True)
class CsvFormat:
    encoding: str
    delimiter: str


def _decode_sample(sample: bytes) -> tuple[str, str]:
    # Incremental decoder with final=False tolerates a multi-byte character
    # truncated at the end of the sample, but still rejects truly invalid bytes.
    try:
        decoder = codecs.getincrementaldecoder("utf-8-sig")()
        return decoder.decode(sample, final=False), "utf-8-sig"
    except UnicodeDecodeError:
        # latin-1 accepts any byte sequence, so it is a safe fallback.
        return sample.decode("latin-1"), "latin-1"


def sniff_csv_format(stream: BinaryIO) -> CsvFormat:
    """Guess encoding and delimiter from the first bytes of the file."""
    sample = stream.read(SAMPLE_BYTES)
    stream.seek(0)
    text, encoding = _decode_sample(sample)
    try:
        delimiter = csv.Sniffer().sniff(text, delimiters=CANDIDATE_DELIMITERS).delimiter
    except csv.Error:
        delimiter = ","
    return CsvFormat(encoding=encoding, delimiter=delimiter)


def read_csv_chunks(
    stream: BinaryIO, chunk_size: int
) -> tuple[list[str], Iterator[pd.DataFrame]]:
    """Return (header, iterator of all-string chunks) for a CSV upload."""
    fmt = sniff_csv_format(stream)
    try:
        # dtype=str + keep_default_na=False: prevent pandas from guessing,
        # type inference (and "NA"/"null" handling) is done manually.
        reader = pd.read_csv(
            stream,
            sep=fmt.delimiter,
            encoding=fmt.encoding,
            dtype=str,
            keep_default_na=False,
            chunksize=chunk_size,
        )
        # Pull the first chunk so header/format errors surface here.
        first = next(reader, None)
    except (pd.errors.ParserError, pd.errors.EmptyDataError, UnicodeDecodeError) as exc:
        raise FileReadError(f"Cannot parse CSV: {exc}") from exc
    if first is None or first.empty:
        raise FileReadError("The file contains no data rows")
    header = list(first.columns)

    def _iter() -> Iterator[pd.DataFrame]:
        yield first
        yield from reader

    return header, _iter()


def read_xlsx_chunks(
    stream: BinaryIO, chunk_size: int
) -> tuple[list[str], Iterator[pd.DataFrame]]:
    """Return (header, chunks) for the FIRST sheet of an XLSX upload.

    openpyxl in read_only mode streams rows instead of loading the whole
    workbook, but it remains the slowest path: documented as a known limit.
    """
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many unrelated exception types
        raise FileReadError(f"Cannot parse XLSX: {exc}") from exc

    rows = wb.worksheets[0].iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        raise FileReadError("The first sheet is empty")
    header = [str(h) if h is not None else "" for h in header_row]
    first_row = next(rows, None)
    if first_row is None:
        raise FileReadError("The file contains no data rows")

    def _iter() -> Iterator[pd.DataFrame]:
        buffer: list[list[str]] = [["" if v is None else str(v) for v in first_row]]
        for row in rows:
            buffer.append(["" if v is None else str(v) for v in row])
            if len(buffer) >= chunk_size:
                yield pd.DataFrame(buffer, columns=header)
                buffer = []
        if buffer:
            yield pd.DataFrame(buffer, columns=header)
        wb.close()

    return header, _iter()


def read_chunks(
    stream: BinaryIO, filename: str, chunk_size: int
) -> tuple[list[str], Iterator[pd.DataFrame]]:
    """Dispatch on file extension. Single entry point used by the services."""
    name = filename.lower()
    if name.endswith((".csv", ".txt")):  # allowing .txt with .csv formatting
        return read_csv_chunks(stream, chunk_size)
    if name.endswith(".xlsx"):
        return read_xlsx_chunks(stream, chunk_size)
    raise FileReadError("Unsupported file type: expected .csv or .xlsx")
